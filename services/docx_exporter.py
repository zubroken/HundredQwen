"""
搜索结果导出为 Word (.docx)
"""
from __future__ import annotations
import os
import tempfile
from typing import List
import logging

logger = logging.getLogger(__name__)


class DocxExporter:
    """知识库搜索结果 → .docx 文件"""

    def export_entry(self, entry: dict) -> str:
        """单条导出，返回文件路径"""
        from docx import Document
        from docx.shared import Pt, Inches
        from docx.enum.text import WD_ALIGN_PARAGRAPH

        doc = Document()

        # 标题
        title = doc.add_heading(entry.get("title", "知识库条目"), level=1)
        title.alignment = WD_ALIGN_PARAGRAPH.LEFT

        # 元信息
        meta = doc.add_paragraph()
        meta.style = doc.styles['Normal']
        meta_run = meta.add_run(f"来源: {entry.get('source', entry.get('doc_title', '知识库'))}")
        meta_run.font.size = Pt(10)
        meta_run.font.color.rgb = None  # 默认黑色

        if entry.get("authors"):
            meta.add_run(f"\n作者: {', '.join(entry['authors']) if isinstance(entry['authors'], list) else entry['authors']}")

        if entry.get("published") or entry.get("page"):
            page_info = entry.get("published") or f"第{entry.get('page', '')}页"
            meta.add_run(f"\n日期/页码: {page_info}")

        doc.add_paragraph("")  # 空行

        # 正文
        text = entry.get("full_text", entry.get("text", ""))
        text_para = doc.add_paragraph(text)
        text_para.style = doc.styles['Normal']

        # 关键词
        keywords = entry.get("keywords", [])
        if keywords:
            doc.add_paragraph("")
            kw_para = doc.add_paragraph()
            kw_run = kw_para.add_run(f"关键词: {', '.join(keywords[:10])}")
            kw_run.font.size = Pt(9)

        # 保存
        safe_title = "".join(c for c in entry.get("title", "export") if c.isalnum() or c in " _-")[:40]
        filepath = os.path.join(tempfile.gettempdir(), f"{safe_title}.docx")
        doc.save(filepath)
        logger.info(f"导出 docx: {filepath}")
        return filepath

    def export_batch(self, entries: List[dict]) -> str:
        """多条合并导出"""
        from docx import Document
        from docx.shared import Pt
        from docx.enum.text import WD_ALIGN_PARAGRAPH

        doc = Document()
        doc.add_heading("知识库搜索结果", level=0)

        for i, entry in enumerate(entries, 1):
            doc.add_heading(f"{i}. {entry.get('title', '未命名')}", level=2)

            meta = doc.add_paragraph()
            meta_run = meta.add_run(f"来源: {entry.get('source', entry.get('doc_title', '知识库'))}")
            meta_run.font.size = Pt(9)

            if entry.get("authors"):
                authors = entry["authors"]
                if isinstance(authors, list):
                    meta.add_run(f"\n作者: {', '.join(authors)}")
                else:
                    meta.add_run(f"\n作者: {authors}")

            if entry.get("published"):
                meta.add_run(f"\n发布日期: {entry['published']}")

            text = entry.get("full_text", entry.get("text", ""))
            doc.add_paragraph(text[:1500])

            if i < len(entries):
                doc.add_paragraph("─" * 40)

        filepath = os.path.join(tempfile.gettempdir(), "知识库搜索结果.docx")
        doc.save(filepath)
        return filepath
